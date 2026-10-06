"""Імпорт готових ТЗ (xlsx — наші Advanced/Base або довільні в стилі «H1/H2/H3 + коментарі + ключі»; docx; текст)."""
import io
import re

import openpyxl

HEAD_RE = re.compile(r"^\s*(?:\d+\.\s*)?(H[1-3])\s*[:.\-–]?\s*(.+)$", re.I)
SKIP_SHEETS = ("зведення", "методика", "аналіз топу", "свод", "сводка")
FREQ_RE = re.compile(r"^\s*\d+\s*([–\-]\s*\d+)?(\s*\(.*\))?\s*$")
_N = r"(\d{1,2}[ \u00a0]\d{3}|\d+)"
VOL_RE = re.compile(r"(?<![\w])" + _N + r"\s*[–\-]\s*" + _N + r"\s*(?:зн|знак|символ)", re.I)
META = {"title": "title", "description": "desc", "рекомендований url": "url"}


def _lang_of(text: str, sheet: str = "") -> str:
    s = sheet.lower()
    if re.search(r"\b(ru|рус)\b", s):
        return "ru"
    if re.search(r"\b(ua|uk|укр)\b", s):
        return "ua"
    ua = len(re.findall(r"[іїєґ]", text.lower()))
    ru = len(re.findall(r"[ыэъё]", text.lower()))
    return "ru" if ru > ua else "ua"


def _num(s: str) -> int:
    return int(re.sub(r"\D", "", s) or 0)


def estimate_volume(blocks) -> str:
    """Сума обсягів з коментарів (з урахуванням «кожен H3»)."""
    lo = hi = 0
    for i, b in enumerate(blocks):
        if b[0] != "c":
            continue
        txt = b[1]
        m = VOL_RE.search(txt)
        if not m:
            continue
        a, z = _num(m.group(1)), _num(m.group(2))
        if re.search(r"кожен\s+H3|каждый\s+H3|кожного\s+H3", txt, re.I):
            n = 0
            for nb in blocks[i + 1:]:
                if nb[0] == "h2":
                    break
                if nb[0] == "h3":
                    n += 1
            intro = re.search(r"вступ[^\d]{0,15}(\d[\d\s]{1,5})", txt, re.I)
            base = _num(intro.group(1)) if intro else 0
            lo += base + a * max(1, n)
            hi += base + z * max(1, n)
        else:
            lo += a
            hi += z
    if not lo:
        return ""
    r = lambda x: f"{round(x, -2):,}".replace(",", " ")  # noqa: E731
    return f"{r(lo)}–{r(hi)} зн. з пробілами (оцінка за коментарями ТЗ)"


def _is_red(cell) -> bool:
    try:
        rgb = cell.font.color.rgb if cell.font and cell.font.color else None
        return isinstance(rgb, str) and rgb.upper().endswith("FF0000")
    except Exception:  # noqa: BLE001
        return False


def parse_sheet(ws) -> dict | None:
    rows = list(ws.iter_rows())
    blocks, meta, volume = [], {}, ""
    after_faq = False
    for row in rows:
        c = row[0] if row else None
        v = c.value if c is not None else None
        if v is None or str(v).strip() == "":
            continue
        s = str(v).strip()
        low = s.lower()
        mk = next((k for k in META if low.startswith(k + ":")), None)
        if mk:
            meta[META[mk]] = s.split(":", 1)[1].strip()
            continue
        if low.startswith(("загальний обсяг тексту", "объём текста", "объем текста", "обсяг тексту")):
            volume = s.split(":", 1)[1].strip() if ":" in s else s
            continue
        if low.startswith(("тз на seo-текст", "мова версії", "шаблон сторінки", "інтент видачі", "аналіз топу",
                           "конкуренти-орієнтири", "ризики", "загальні вимоги")) or s.startswith("• "):
            continue
        m = HEAD_RE.match(s)
        if m:
            lvl, text = m.group(1).lower(), m.group(2).strip()
            blocks.append((lvl, text, text))
            after_faq = False
            continue
        if low in ("cta", "cta-блок") or low.startswith("cta"):
            blocks.append(("cta",))
            continue
        if low in ("faq", "часті запитання", "часто задаваемые вопросы", "частые вопросы"):
            blocks.append(("faq",))
            after_faq = True
            continue
        if after_faq and not _is_red(c) and s.endswith("?"):
            blocks.append(("q", s, s))
            continue
        blocks.append(("c", s))
    if not any(b[0] == "h1" for b in blocks) and sum(1 for b in blocks if b[0] == "h2") < 2:
        return None
    # ключі: шукаємо пару колонок «фраза | частота»
    kw, lsi = [], []
    target = kw
    for col in range(1, min(ws.max_column, 8)):
        pairs = []
        for row in rows:
            if len(row) <= col + 1:
                continue
            k, f = row[col].value, row[col + 1].value
            if k is None:
                continue
            ks = str(k).strip()
            kl = ks.lower()
            if "lsi" in kl:
                target = lsi
                continue
            if kl.startswith(("keyword", "ключов", "ключі", "семантика", "сумарна")) or kl.startswith("діапазони"):
                if kl.startswith("семантика"):
                    break
                continue
            if f is not None and FREQ_RE.match(str(f)):
                target.append((ks, str(f).strip()))
                pairs.append(1)
        if pairs:
            break
    if not volume:
        volume = estimate_volume(blocks)
    h1 = next((b[1] for b in blocks if b[0] == "h1"), ws.title)
    lang = _lang_of(" ".join(b[1] for b in blocks if b[0] in ("h1", "h2", "h3", "q")), ws.title)
    tz = {"short": ws.title[:20], "name_ru": h1, "name_ua": h1, "template": "", "intent": "", "volume": volume,
          "top_note": "", "url_ru": meta.get("url", ""), "url_ua": meta.get("url", ""),
          f"title_{lang}": meta.get("title", ""), f"desc_{lang}": meta.get("desc", ""), "risks": "",
          "blocks": blocks, "kw_ru": [], "kw_ua": [], "lsi_ru": [], "lsi_ua": [], "competitors": []}
    tz[f"kw_{lang}"], tz[f"lsi_{lang}"] = kw, lsi
    return {"sheet": ws.title, "lang": lang, "name": h1, "tz": tz,
            "stats": {"headings": sum(1 for b in blocks if b[0] in ("h1", "h2", "h3")), "keywords": len(kw) + len(lsi),
                      "volume": volume or "не вказано"}}


def sheet_text(ws) -> str:
    lines = []
    for row in ws.iter_rows():
        vals = []
        for c in row:
            if c.value is None or str(c.value).strip() == "":
                continue
            v = str(c.value).strip()
            vals.append(f"[коментар] {v}" if _is_red(c) else v)
        if vals:
            lines.append(" | ".join(vals))
    return "\n".join(lines)


def parse_xlsx(data: bytes) -> list[dict]:
    """Розпізнані аркуші → готові ТЗ; нерозпізнані непорожні → raw (розбере Claude)."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    out = []
    for ws in wb.worksheets:
        if ws.title.strip().lower().startswith(SKIP_SHEETS):
            continue
        r = parse_sheet(ws)
        if r:
            out.append(r)
        else:
            raw = sheet_text(ws)
            if len(raw) > 300 and re.search(r"\bH[1-3]\b|заголов|структур", raw, re.I):
                out.append(raw_entry(raw, ws.title))
    return out


def raw_entry(raw: str, title: str) -> dict:
    return {"sheet": title, "lang": _lang_of(raw), "name": title, "tz": None, "raw": raw,
            "stats": {"headings": "розбере Claude", "keywords": "—", "volume": "—"}}


def docx_text(data: bytes) -> str:
    from docx import Document
    d = Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        st = (p.style.name or "").lower()
        if st.startswith("heading") and st[-1:].isdigit() and not HEAD_RE.match(t):
            t = f"H{st[-1]}: {t}"
        lines.append(t)
    for tbl in d.tables:
        for row in tbl.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


IMPORT_SYS = """Ти — SEO-спеціаліст. Тобі дали готове ТЗ копірайтеру у довільному форматі. Перенеси його в структуру submit_tz БЕЗ змін змісту:
- blocks: заголовки H1/H2/H3 дослівно (і в ru, і в ua пиши однаковий текст — мовою оригіналу); усі вимоги/коментарі до блоку — окремими блоками type "c" (дослівно або дуже близько); CTA → {"type":"cta"} + коментар; FAQ → {"type":"faq"} + питання type "q".
- volume — загальний обсяг з ТЗ; якщо не вказано — порахуй суму обсягів блоків і познач «(оцінка)».
- kw_* та lsi_* — ключі з частотою для мови ТЗ (визнач мову: ua або ru); іншу мову залиш порожньою.
- Нічого не вигадуй і не додавай від себе. Поля intent/top_note/competitors залиш порожніми, якщо в ТЗ їх немає."""


def llm_import(raw: str, title: str) -> dict:
    from . import llm
    from .prompts import TOOL
    data = llm.call_tool(IMPORT_SYS, f"ТЗ «{title}»:\n\n{raw[:60000]}", TOOL, 16000, check=lambda d: bool(d.get("blocks")))
    tz = llm.normalize(data)
    lang = "ru" if tz.get("kw_ru") and not tz.get("kw_ua") else _lang_of(" ".join(b[1] for b in tz["blocks"] if b[0] in ("h1", "h2", "h3", "q")))
    return {"sheet": title, "lang": lang, "name": tz.get(f"name_{lang}") or title, "tz": tz,
            "stats": {"headings": sum(1 for b in tz["blocks"] if b[0] in ("h1", "h2", "h3")),
                      "keywords": len(tz.get(f"kw_{lang}", [])) + len(tz.get(f"lsi_{lang}", [])), "volume": tz.get("volume") or "не вказано"}}
