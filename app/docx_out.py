"""Тексти у Word: кожна сторінка × мова — окремий розділ з нової сторінки + звіт перевірки."""
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor


def _link(par, text, url, bold=False):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1155CC"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if bold:
        rpr.append(OxmlElement("w:b"))
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r)
    par._p.append(h)


def _runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        if bold:
            part = part[2:-2]
        for j, seg in enumerate(re.split(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", part)):
            k = j % 3
            if k == 0 and seg:
                par.add_run(seg).bold = bold or None
            elif k == 1:
                label = seg
            elif k == 2:
                _link(par, label, seg, bold)


def _md(doc, md, justify=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    J = AL.JUSTIFY if justify else None

    def para(style=None):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        if J is not None:
            p.alignment = J
        return p

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        m = re.match(r"^(#{1,3})\s+(.+)$", ln)
        if m:
            h = doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
            if justify:
                h.alignment = AL.LEFT
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|?\s*:?-{2,}", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            i -= 1
            if rows:
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=cols)
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci, v in enumerate(r):
                        cp = t.cell(ri, ci).paragraphs[0]
                        v = v.replace("**", "")
                        _runs(cp, f"**{v}**" if ri == 0 and v else v)
                        if justify:
                            cp.alignment = AL.CENTER if ri == 0 else AL.JUSTIFY
        elif re.match(r"^\s*[-*•]\s+", ln):
            _runs(para("List Bullet"), re.sub(r"^\s*[-*•]\s+", "", ln))
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            _runs(para("List Number"), re.sub(r"^\s*\d+[.)]\s+", "", ln))
        elif ln.strip():
            _runs(para(), ln.strip())
        i += 1


def _base_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    return doc


def build_article(t: dict, path):
    """Чиста стаття для медіа: файл починається одразу з H1, без змісту, мета й службових блоків."""
    doc = _base_doc()
    _md(doc, t["markdown"], justify=True)
    doc.save(path)


def _slug(s: str) -> str:
    s = re.sub(r"[^\w\s-]", "", (s or "").lower(), flags=re.U)
    return re.sub(r"[\s_-]+", "-", s).strip("-")[:60] or "stattia"


def build_media(texts: list[dict], outdir, slug: str) -> list[str]:
    """Медіа: окремий .docx на кожну статтю × мову + технічний Excel; усе разом у zip. Повертає імена файлів."""
    import zipfile
    import openpyxl
    from openpyxl.styles import Alignment, Font
    names = []
    for t in texts:
        h1 = next((m.group(1) for m in re.finditer(r"^#\s+(.+)$", t["markdown"], re.M)), t["title"])
        fn = f"{t['id']}_{t['lang']}_{_slug(h1)}.docx"
        build_article(t, outdir / fn)
        t["file"] = fn
        names.append(fn)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Статті"
    head = ["#", "Мова", "Файл", "H1", "Title", "Description", "Теги", "Обсяг, зн.", "Стиль", "H3 → список (блоки)",
            "Прибрано з ТЗ", "Потребує уваги", "Рекомендації"]
    ws.append(head)
    for t in texts:
        rep, meta = t["report"], t.get("meta") or {}
        h1 = next((m.group(1) for m in re.finditer(r"^#\s+(.+)$", t["markdown"], re.M)), t["title"])
        ws.append([t["id"], t["lang"].upper(), t.get("file", ""), h1, meta.get("title") or t.get("title_tag") or "",
                   meta.get("description") or t.get("desc") or "", ", ".join(meta.get("tags") or []), rep["chars"],
                   t.get("style_name", ""), "; ".join(rep.get("converted") or []), "; ".join(rep.get("dropped") or []),
                   "\n".join(rep["issues"]) or "—", "\n".join(rep.get("warnings") or [])])
    ks = wb.create_sheet("Ключі")
    ks.append(["#", "Мова", "Тип", "Ключ / LSI", "Входжень", "Норма"])
    for t in texts:
        for kind, k, n, f in t["report"].get("keywords") or []:
            ks.append([t["id"], t["lang"].upper(), kind, k, n, str(f)])
    fs = wb.create_sheet("Джерела")
    fs.append(["#", "Мова", "Факт", "У тексті", "Організація", "Рік", "Географія", "URL", "Цитата"])
    for t in texts:
        used = set(t["report"].get("facts_used") or [])
        for f in t.get("facts") or []:
            fs.append([t["id"], t["lang"].upper(), f["id"], "так" if f["id"] in used else "—", f.get("org"), f.get("year"),
                       f.get("geo"), f.get("url"), f.get("quote")])
    ls = wb.create_sheet("Мовні зауваження")
    ls.append(["#", "Мова", "Зауваження LanguageTool (перевірте вручну)"])
    for t in texts:
        for x in t.get("spelling") or []:
            ls.append([t["id"], t["lang"].upper(), x])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = 22
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    tech = f"Tech_{slug}.xlsx"
    wb.save(outdir / tech)
    zname = f"Articles_{slug}.zip"
    with zipfile.ZipFile(outdir / zname, "w", zipfile.ZIP_DEFLATED) as z:
        for fn in names + [tech]:
            z.write(outdir / fn, fn)
    return [zname, tech]


def build(texts: list[dict], path):
    """texts: [{title, lang, url, title_tag, desc, markdown, report, spelling}]"""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    doc.add_heading("SEO-тексти", 0)
    p = doc.add_paragraph("Зміст: ")
    for t in texts:
        doc.add_paragraph(f"{t['title']} — {t['lang'].upper()} ({t['report']['chars']} зн.)" + (f", {t['text_type_name'][0].lower() + t['text_type_name'][1:]}" if t.get("text_type") and t["text_type"] != "landing" else "") + (f", стиль: {t['style_name']}" if t.get("style_name") else ""), style="List Bullet")
    for t in texts:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        meta = doc.add_paragraph()
        r = meta.add_run(f"{t['title']} · {t['lang'].upper()} · {t.get('url') or ''}")
        r.italic = True
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if t.get("title_tag"):
            mp = doc.add_paragraph(); mp.add_run("Title: ").bold = True; mp.add_run(t["title_tag"])
        if t.get("desc"):
            mp = doc.add_paragraph(); mp.add_run("Description: ").bold = True; mp.add_run(t["desc"])
        _md(doc, t["markdown"], justify=t.get("text_type") in ("article", "media"))
        rep = t["report"]
        doc.add_heading("Перевірка тексту (службовий блок, не публікувати)", 3)
        doc.add_paragraph(f"Обсяг: {rep['chars']} зн. з пробілами. Раундів редагування: {t.get('rounds', 0)}."
                          + (f" Тип: {t['text_type_name']}." if t.get("text_type_name") else "")
                          + (f" Стиль: {t['style_name']}." if t.get("style_name") else ""))
        fs = rep.get("format") or {}
        if fs:
            doc.add_paragraph(f"Формат: {fs.get('paragraphs')} абзаців, у середньому {fs.get('avg_sentences')} речення / "
                              f"{fs.get('avg_words')} слів на абзац; списків {fs.get('lists')} ({fs.get('list_share')}% тексту).")
        if rep["issues"]:
            doc.add_paragraph().add_run("Потребує уваги (не вдалося виправити автоматично):").bold = True
            for x in rep["issues"]:
                doc.add_paragraph(x, style="List Bullet")
        else:
            doc.add_paragraph().add_run("Критичних зауважень немає.").bold = True
        if rep.get("warnings"):
            doc.add_paragraph("Рекомендації (не критично):")
            for x in rep["warnings"]:
                doc.add_paragraph(x, style="List Bullet")
        secs = [s for s in rep.get("sections", []) if s[1] and s[2]]
        if secs:
            doc.add_paragraph("Обсяги розділів:")
            for t_, a, lo, hi in secs:
                doc.add_paragraph(f"{t_}: {a} зн. (ціль {lo}–{hi})", style="List Bullet")
        if t.get("spelling"):
            doc.add_paragraph("Залишкові зауваження LanguageTool (перевірте вручну, частина може бути хибною):")
            for x in t["spelling"][:15]:
                doc.add_paragraph(x, style="List Bullet")
        if rep.get("placeholders"):
            doc.add_paragraph("Плейсхолдери: " + "; ".join(dict.fromkeys(rep["placeholders"])))
        if t.get("facts"):
            used = set(rep.get("facts_used") or [])
            doc.add_paragraph().add_run(f"Факти з джерелами (перевірено дослівно на сторінці джерела; використано {len(used)} з {len(t['facts'])}):").bold = True
            ft = doc.add_table(rows=1, cols=4)
            ft.style = "Table Grid"
            for i, h in enumerate(["#", "Факт", "Джерело", "У тексті"]):
                ft.rows[0].cells[i].text = h
            for f in t["facts"]:
                row = ft.add_row().cells
                row[0].text = f["id"]
                row[1].text = f.get("claim_" + ("ru" if t["lang"] == "ru" else "ua"), "") + f"\nЦитата: «{f.get('quote', '')}»"
                row[2].text = f"{f.get('org')}, {f.get('year')}, {f.get('geo')}"
                _link(row[2].add_paragraph(), f["url"], f["url"])
                row[3].text = "так" if f["id"] in used else "—"
        if rep.get("keywords"):
            tbl = doc.add_table(rows=1, cols=4)
            tbl.style = "Table Grid"
            for i, h in enumerate(["Тип", "Ключ / LSI", "Входжень", "Норма"]):
                tbl.rows[0].cells[i].text = h
            for kind, k, n, f in rep["keywords"]:
                row = tbl.add_row().cells
                row[0].text, row[1].text, row[2].text, row[3].text = kind, k, str(n), str(f)
    doc.save(path)
