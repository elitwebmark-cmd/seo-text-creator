"""Автоматичний збір семантики: послуга або URL → сиди → розширення → частотність → чистка й кластеризація."""
import asyncio
import json
import re

from . import config, llm, scraper, serp, volumes

S_SYS = """Ти — senior SEO-спеціаліст, який збирає семантику для комерційної сторінки послуги на сайті в Україні.
За назвою послуги або вмістом сторінки сформуй сиди — базові пошукові запити, якими українці шукають цю послугу в Google.
Правила:
- Окремо для української (ua) і російської (ru) мови. Запити пишемо так, як їх реально вводять у пошук: у нижньому регістрі, без розділових знаків, з типовими сленговими й транслітерованими варіантами (наприклад «serm», «сео», «ютуб»).
- 10–15 сидів на мову: ядро послуги, синоніми, комерційні модифікатори («послуги», «замовити», «ціна», «під ключ», «агенція»), 2–3 інформаційних запити, що можуть потрапити на сторінку послуги («що таке …»).
- Мінус-слова: те, що точно не наш інтент (вакансії, курси, безкоштовно, скачати, своїми руками тощо — залежно від ніші).
- Не вигадуй брендів."""

S_TOOL = {
    "name": "submit_seeds",
    "description": "Сиди для збору семантики",
    "input_schema": {
        "type": "object",
        "properties": {
            "page_name_ua": {"type": "string"},
            "page_name_ru": {"type": "string"},
            "service_summary": {"type": "string", "description": "1–2 речення: що саме продаємо на сторінці"},
            "seeds_ua": {"type": "array", "items": {"type": "string"}},
            "seeds_ru": {"type": "array", "items": {"type": "string"}},
            "negative": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["page_name_ua", "page_name_ru", "service_summary", "seeds_ua", "seeds_ru", "negative"],
    },
}

C_SYS = """Ти — senior SEO-спеціаліст. Тобі дано кандидатів у семантичне ядро для однієї комерційної сторінки послуги (з частотністю Google Ads по Україні, якщо вона є).
Завдання: почистити й розкласти ключі на сторінки.
1. main — ключі, які веде саме ця сторінка: той самий інтент (замовити/отримати послугу, дізнатися ціну, обрати виконавця) + близькі інформаційні запити на кшталт «що таке …», які в Google ранжуються на сторінках послуг. Типово 5–25 ключів на мову.
2. extra_pages — якщо серед кандидатів є окремі інтенти з помітною частотністю, які краще вести іншою сторінкою (інша послуга, B2B-варіант, місто, статті-гайди), запропонуй їх як окремі кластери з назвою та ключами. Не більше 4.
3. Відкинь: нерелевантне, мінус-слова, вакансії/курси/безкоштовне, чужі бренди, дублі, запити іншою мовою (ua-ключі мають бути українською або латиницею, ru — російською або латиницею), граматично биті фрази.
4. Ключі з нульовою або невідомою частотністю лиши, лише якщо вони точно релевантні й природні (≤ 5 таких на мову).
5. Не змінюй написання ключів і не вигадуй нових — бери тільки з кандидатів.
6. main_keyword — головний ключ мови (найчастотніший точний за інтентом)."""

_KWS = {"type": "array", "items": {"type": "string"}}
C_TOOL = {
    "name": "submit_semantics",
    "description": "Почищена семантика",
    "input_schema": {
        "type": "object",
        "properties": {
            "main_ua": _KWS, "main_ru": _KWS,
            "main_keyword_ua": {"type": "string"}, "main_keyword_ru": {"type": "string"},
            "extra_pages": {
                "type": "array",
                "items": {"type": "object", "properties": {
                    "name_ua": {"type": "string"}, "name_ru": {"type": "string"},
                    "keywords_ua": _KWS, "keywords_ru": _KWS, "why": {"type": "string"}},
                    "required": ["name_ua", "name_ru", "keywords_ua", "keywords_ru"]},
            },
            "notes": {"type": "string", "description": "Коротко: що відкинуто і чому, на що звернути увагу"},
        },
        "required": ["main_ua", "main_ru", "extra_pages"],
    },
}


def _is_url(s: str) -> bool:
    return bool(re.match(r"^https?://", s.strip()))


def _norm(k: str) -> str:
    return " ".join(k.lower().replace("ё", "е").split())


async def collect_one(entry: str, langs: list[str], project: dict, log) -> dict:
    entry = entry.strip()
    context = f"Послуга: {entry}"
    if _is_url(entry):
        log(f"читаю сторінку {entry}")
        pages = await scraper.fetch_many([entry])
        p = pages[0]
        if p.get("error"):
            context = f"Сторінка {entry} (не вдалося завантажити: {p['error']}). Визнач послугу з URL."
        else:
            hs = "\n".join(f"{l}: {t}" for l, t in p.get("headings", [])[:40])
            context = (f"Приклад сторінки: {entry}\nTitle: {p.get('title')}\nDescription: {p.get('description')}\n"
                       f"Заголовки:\n{hs}\nФрагмент тексту:\n{p.get('text', '')[:2500]}")
    if project.get("notes"):
        context += f"\nДодатково: {project['notes']}"

    if config.MOCK:
        seeds = {"page_name_ua": entry[:60], "page_name_ru": entry[:60], "service_summary": entry,
                 "seeds_ua": [entry.lower(), f"{entry.lower()} послуги"], "seeds_ru": [entry.lower(), f"{entry.lower()} услуги"],
                 "negative": ["вакансії"]}
    else:
        log("Claude формує сиди")
        seeds = await asyncio.to_thread(llm.call_tool, S_SYS, context, S_TOOL, 4000)

    neg = [n.lower() for n in seeds.get("negative", []) if n]
    cand: dict[str, dict[str, int | None]] = {}
    questions: dict[str, list[str]] = {}
    for lang in langs:
        sd = [s for s in seeds.get(f"seeds_{lang}", []) if s][:15]
        pool: dict[str, int | None] = {_norm(s): None for s in sd}
        log(f"{lang.upper()}: підказки Google для {len(sd)} сидів")
        auto = await asyncio.gather(*(volumes.autocomplete(s, lang) for s in sd[:12]))
        for lst in auto:
            for k in lst:
                pool.setdefault(_norm(k), None)
        qs = []
        for s in sd[:2]:
            try:
                res = await serp.search(s, lang)
                for k in res.get("related", []):
                    pool.setdefault(_norm(k), None)
                qs += res.get("paa", [])
            except Exception as e:  # noqa: BLE001
                log(f"видача за «{s}»: {e}")
        questions[lang] = list(dict.fromkeys(qs))
        if not (volumes.dfs_enabled() or config.MOCK):
            log("DataForSEO вимкнено (не задано логін/пароль) — частотність не збирається")
        if (volumes.dfs_enabled() and config.DATAFORSEO_IDEAS) or config.MOCK:
            try:
                log(f"{lang.upper()}: ідеї Google Ads")
                for k, v in (await volumes.keyword_ideas(sd, lang)).items():
                    pool[_norm(k)] = v
            except Exception as e:  # noqa: BLE001
                log(f"ідеї Google Ads: {e}")
        pool = {k: v for k, v in pool.items() if k and not any(n in k for n in neg)}
        missing = [k for k, v in pool.items() if v is None]
        if missing and (volumes.dfs_enabled() or config.MOCK):
            try:
                log(f"{lang.upper()}: частотність для {len(missing)} ключів")
                vols = await volumes.search_volume(missing, lang)
                log(f"{lang.upper()}: DataForSEO повернув частотність для {sum(1 for v in vols.values() if v)} з {len(missing)} ключів")
                for k in missing:
                    if k in vols:
                        pool[k] = vols[k]
            except Exception as e:  # noqa: BLE001
                log(f"частотність: {e}")
        ranked = sorted(pool.items(), key=lambda x: -(x[1] or 0))[:350]
        cand[lang] = dict(ranked)
        log(f"{lang.upper()}: кандидатів {len(cand[lang])}")

    if config.MOCK:
        sem = {f"main_{l}": list(cand.get(l, {}))[:8] for l in ("ua", "ru")}
        sem["extra_pages"] = []
    else:
        lines = [f"Послуга: {seeds.get('service_summary')}", f"Мінус-слова: {', '.join(neg) or '—'}"]
        for lang in langs:
            lines.append(f"\n## Кандидати {lang.upper()} (ключ — частотність, «?» = невідома)")
            lines += [f"{k} — {'?' if v is None else v}" for k, v in cand[lang].items()]
        log("Claude чистить і кластеризує")
        sem = await asyncio.to_thread(llm.call_tool, C_SYS, "\n".join(lines), C_TOOL, 8000)

    def pack(keys, lang):
        pool = cand.get(lang, {})
        out = []
        for k in keys or []:
            nk = _norm(k)
            if nk in pool and all(nk != x[0] for x in out):
                out.append((nk, None if pool[nk] is None else int(pool[nk])))
        return sorted(out, key=lambda x: -(x[1] or 0))

    result = {"entry": entry, "summary": seeds.get("service_summary", ""), "notes": sem.get("notes", ""),
              "pages": [{"name_ua": seeds.get("page_name_ua", entry), "name_ru": seeds.get("page_name_ru", entry),
                         "ua": pack(sem.get("main_ua"), "ua") if "ua" in langs else [],
                         "ru": pack(sem.get("main_ru"), "ru") if "ru" in langs else [],
                         "main": True}],
              "questions": questions, "volumes_source": "Google Ads (DataForSEO)" if volumes.dfs_enabled() else "немає (не задано DataForSEO)"}
    for ep in sem.get("extra_pages", []) or []:
        pg = {"name_ua": ep.get("name_ua", ""), "name_ru": ep.get("name_ru", ""),
              "ua": pack(ep.get("keywords_ua"), "ua") if "ua" in langs else [],
              "ru": pack(ep.get("keywords_ru"), "ru") if "ru" in langs else [],
              "main": False, "why": ep.get("why", "")}
        if pg["ua"] or pg["ru"]:
            result["pages"].append(pg)
    return result


def to_semantics(results: list[dict]):
    """Формат, який очікує інтерфейс: кластери ua/ru + пари."""
    ua, ru, pairs = [], [], []
    for r in results:
        for p in r["pages"]:
            iu = ir = None
            if p["ua"]:
                ua.append({"name": p["name_ua"] or p["name_ru"], "keywords": [list(x) for x in p["ua"]]})
                iu = len(ua) - 1
            if p["ru"]:
                ru.append({"name": p["name_ru"] or p["name_ua"], "keywords": [list(x) for x in p["ru"]]})
                ir = len(ru) - 1
            pairs.append({"ua": iu, "ru": ir, "main": p.get("main", True), "why": p.get("why", "")})
    return {"ua": ua, "ru": ru, "pages": pairs}


def semantics_xlsx(results: list[dict], path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for lang in ("ua", "ru"):
        ws = wb.create_sheet(lang)
        ws.append([None, "кластер", "Ключове слово", "Частотність", "Мова"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in results:
            for p in r["pages"]:
                if not p[lang]:
                    continue
                ws.append([None, p[f"name_{lang}"] + ("" if p.get("main", True) else " (дод. сторінка)")])
                ws.cell(ws.max_row, 2).font = Font(bold=True)
                for k, v in p[lang]:
                    ws.append([None, None, k, v, lang.upper()])
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 48
        q = wb.create_sheet(f"питання {lang}")
        q.append(["Сторінка", "Питання з «Люди також питають»"])
        for r in results:
            for x in r.get("questions", {}).get(lang, []):
                q.append([r["entry"], x])
    wb.save(path)


async def run(entries: list[str], langs: list[str], project: dict, log) -> list[dict]:
    out = []
    for e in entries:
        try:
            out.append(await collect_one(e, langs, project, lambda m, e=e: log(f"[{e[:40]}] {m}")))
        except Exception as ex:  # noqa: BLE001
            log(f"[{e[:40]}] ПОМИЛКА: {ex}")
    return out


def dumps(results):
    return json.dumps(results, ensure_ascii=False, indent=1)
