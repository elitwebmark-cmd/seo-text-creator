"""Збірка Excel: Advanced (детальне ТЗ) і Base (як у шаблоні-прикладі)."""
import re
import time

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
f_title = Font(name=F, size=14, bold=True)
f_meta_k = Font(name=F, size=10, bold=True)
f_meta = Font(name=F, size=10)
f_h1 = Font(name=F, size=13, bold=True)
f_h2 = Font(name=F, size=12, bold=True)
f_h3 = Font(name=F, size=11, bold=True)
f_c = Font(name=F, size=10, color="FFFF0000")
f_t = Font(name=F, size=10)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFFFF")
f_note = Font(name=F, size=9, italic=True, color="FF666666")
fill_h1 = PatternFill("solid", fgColor="FFD9EAD3")
fill_h2 = PatternFill("solid", fgColor="FFFFFF00")
fill_hdr = PatternFill("solid", fgColor="FF3C4B64")
fill_meta = PatternFill("solid", fgColor="FFF3F3F3")
fill_sec = PatternFill("solid", fgColor="FFDDE6F5")
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="FFBFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

GENERAL = [
    "Унікальність тексту — від 90% (text.ru / Advego). «Вода» — до 15%, без переспаму.",
    "Основний ключ — у H1, у першому абзаці та в 1–2 H2. Точні входження розподіляти рівномірно, не більше одного ключа в реченні.",
    "Частоти в таблиці ключів — на весь текст, разом із заголовками. «(варіації)» — можна змінювати словоформу, число, відмінок. LSI — діапазон «мін–макс» за ТОПом, перевищувати не можна.",
    "Обсяги блоків — орієнтовні (±15%). Загальний обсяг — обов’язковий.",
    "Форматування: абзаци до 3–4 речень, списки, таблиці там, де вони вказані в ТЗ, нумеровані етапи.",
    "Тон: експертний, від імені компанії («мы» / «ми»), звернення до читача на «вы» / «ви». Без канцеляризмів і загальних фраз без доказів.",
    "Дані з позначкою «УТОЧНИТИ» (ціни, кейси, цифри, імена) не вигадувати — ставити плейсхолдер [уточнити].",
    "Не згадувати конкурентів. Цифри ринку давати без назв компаній.",
    "UA-версія — не калька з RU: пишемо українською з нуля за UA-семантикою.",
    "FAQ: відповідь на кожне питання — 250–400 зн. Під FAQ закласти мікророзмітку FAQPage.",
]


def est_h(text, width_chars=95, line=13):
    lines = max(1, sum(max(1, -(-len(p) // width_chars)) for p in str(text).split("\n")))
    return min(409, max(15, lines * line))


def vol_sum(items):
    return int(sum(v for _, v in items if isinstance(v, (int, float))))


def vol_txt(v):
    return "—" if v is None else v


def safe_title(wb, name):
    name = re.sub(r"[\[\]\*\?/\\:]", " ", name)[:31].strip()
    base, i = name, 2
    while name in wb.sheetnames:
        suf = f" ({i})"
        name = base[: 31 - len(suf)] + suf
        i += 1
    return name


def _langs(item):
    return [l for l in ("ru", "ua") if item["sem"].get(l)]


# ------------------------------------------------------------------ Advanced
def _tz_sheet(wb, item, lang):
    c = item["tz"]
    L = 0 if lang == "ru" else 1
    LANG = lang.upper()
    name = c["name_ru"] if lang == "ru" else c["name_ua"]
    ws = wb.create_sheet(safe_title(wb, f"{item['id']} {LANG} {c['short']}"))
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 95, "B": 3, "C": 42, "D": 24, "E": 3, "F": 40, "G": 13}.items():
        ws.column_dimensions[col].width = w
    ws.cell(1, 1, f"ТЗ на SEO-текст: {name} — {LANG}-версія").font = f_title
    rows = [
        ("Рекомендований URL", c.get(f"url_{lang}") or "—"),
        ("Мова версії", "російська" if lang == "ru" else "українська"),
        ("Шаблон сторінки", c.get("template", "")),
        ("Загальний обсяг тексту", c.get("volume", "")),
        ("Title", c.get(f"title_{lang}", "")),
        ("Description", c.get(f"desc_{lang}", "")),
        ("Інтент видачі", c.get("intent", "")),
        ("Аналіз ТОПу і чим будемо кращими", c.get("top_note", "")),
        ("Конкуренти-орієнтири", " · ".join(c.get("competitors", []))),
    ]
    if c.get("risks"):
        rows.append(("Ризики / перелінковка", c["risks"]))
    r = 2
    for k, v in rows:
        cell = ws.cell(r, 1, f"{k}: {v}")
        cell.font = f_meta; cell.alignment = wrap; cell.fill = fill_meta
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = est_h(f"{k}: {v}", 215, 13)
        r += 1
    r += 1
    r0 = r
    for b in c["blocks"]:
        t = b[0]
        if t in ("h1", "h2", "h3"):
            txt = f"{t.upper()}: {b[1 + L]}"
            cell = ws.cell(r, 1, txt)
            cell.font = {"h1": f_h1, "h2": f_h2, "h3": f_h3}[t]
            if t == "h1": cell.fill = fill_h1
            if t == "h2": cell.fill = fill_h2
            cell.alignment = wrap
        elif t == "c":
            cell = ws.cell(r, 1, b[1]); cell.font = f_c; cell.alignment = wrap
            ws.row_dimensions[r].height = est_h(b[1])
        elif t == "cta":
            cell = ws.cell(r, 1, "CTA"); cell.font = f_h2; cell.fill = fill_h2
        elif t == "faq":
            r += 1
            cell = ws.cell(r, 1, "FAQ"); cell.font = f_h2; cell.fill = fill_h2
            r += 1
            ws.cell(r, 1, "Відповідь на кожне питання — 250–400 зн. Мікророзмітка FAQPage.").font = f_c
        elif t in ("q", "t"):
            cell = ws.cell(r, 1, b[1 + L]); cell.font = f_t; cell.alignment = wrap
        r += 1
        if t == "c":
            r += 1
    r += 1
    cell = ws.cell(r, 1, "Загальні вимоги до тексту"); cell.font = f_h2; cell.fill = fill_sec
    r += 1
    for g in GENERAL:
        cell = ws.cell(r, 1, "• " + g); cell.font = f_t; cell.alignment = wrap
        ws.row_dimensions[r].height = est_h(g)
        r += 1

    def hdr(rr, col, a, b):
        for i, v in enumerate((a, b)):
            x = ws.cell(rr, col + i, v); x.font = f_hdr; x.fill = fill_hdr; x.alignment = wrap; x.border = box

    hdr(r0, 3, "Ключові запити (точні / морф.)", "Рекомендована частота")
    rr = r0 + 1
    for k, f in c.get(f"kw_{lang}", []):
        ws.cell(rr, 3, k).font = Font(name=F, size=10, bold=True)
        ws.cell(rr, 4, f).font = f_t
        for col in (3, 4):
            ws.cell(rr, col).border = box; ws.cell(rr, col).alignment = wrap
        rr += 1
    rr += 1
    hdr(rr, 3, "LSI / тематичні слова (за ТОПом)", "Діапазон на текст")
    rr += 1
    for k, f in c.get(f"lsi_{lang}", []):
        ws.cell(rr, 3, k).font = f_t
        ws.cell(rr, 4, f).font = f_t
        for col in (3, 4):
            ws.cell(rr, col).border = box; ws.cell(rr, col).alignment = wrap
        rr += 1
    rr += 1
    n = ws.cell(rr, 3, "Діапазони LSI розраховані за кількістю входжень у сторінках ТОПу, приведеною до нашого обсягу тексту.")
    n.font = f_note; n.alignment = wrap
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)

    sem = item["sem"].get(lang, [])
    hdr(r0, 6, f"Семантика кластера ({LANG})", "Частотність")
    rr = r0 + 1
    for k, v in sem:
        ws.cell(rr, 6, k).font = f_t; ws.cell(rr, 7, vol_txt(v)).font = f_t
        for col in (6, 7):
            ws.cell(rr, col).border = box
        rr += 1
    ws.cell(rr, 6, "Сумарна частотність").font = f_meta_k
    ws.cell(rr, 7, vol_sum(sem)).font = f_meta_k
    ws.freeze_panes = "A2"
    return ws


def build_advanced(items, project, path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    S = wb.create_sheet("Зведення")
    S.sheet_view.showGridLines = False
    dom = project.get("domain") or "проєкту"
    S["A1"] = f"ТЗ на SEO-тексти для {dom} — зведення"; S["A1"].font = f_title
    S["A2"] = (f"Згенеровано {time.strftime('%d.%m.%Y')}. Видача Google Україна (gl=ua, hl=uk / hl=ru) через Serper, "
               "розбір сторінок ТОПу (H1–H3, обсяг, входження ключів, FAQ, ціни), структура — Claude.")
    S["A2"].font = f_meta; S["A2"].alignment = wrap
    S.merge_cells("A2:K2"); S.row_dimensions[2].height = 30
    heads = ["№", "Сторінка", "Шаблон", "URL (RU / UA)", "Обсяг тексту", "Головний ключ RU", "Σ частотн. RU",
             "Головний ключ UA", "Σ частотн. UA", "Аркуші ТЗ", "Ризики / примітки"]
    for i, h in enumerate(heads, 1):
        x = S.cell(4, i, h); x.font = f_hdr; x.fill = fill_hdr; x.alignment = wrap; x.border = box
    for col, w in zip("ABCDEFGHIJK", [5, 30, 24, 30, 16, 28, 11, 28, 11, 24, 60]):
        S.column_dimensions[col].width = w
    r = 5
    for it in items:
        c = it["tz"]
        sheets = [_tz_sheet(wb, it, l).title for l in _langs(it)]
        vals = [it["id"], c.get("name_ua") or c.get("name_ru"), c.get("template"),
                f"{c.get('url_ru') or '—'}\n{c.get('url_ua') or '—'}", c.get("volume"),
                (c["kw_ru"][0][0] if c.get("kw_ru") else "—"), vol_sum(it["sem"].get("ru", [])),
                (c["kw_ua"][0][0] if c.get("kw_ua") else "—"), vol_sum(it["sem"].get("ua", [])),
                "\n".join(sheets), c.get("risks") or "—"]
        for i, v in enumerate(vals, 1):
            x = S.cell(r, i, v); x.font = f_meta; x.alignment = wrap; x.border = box
        S.row_dimensions[r].height = max(45, est_h(vals[-1], 62, 13))
        r += 1
    S.freeze_panes = "A5"
    r += 1
    S.cell(r, 1, "Як читати ТЗ").font = f_h2
    r += 1
    for k, v, fl in [("H1", "зелена заливка — заголовок сторінки", fill_h1),
                     ("H2 / CTA / FAQ", "жовта заливка — розділи сторінки", fill_h2),
                     ("H3", "жирний без заливки — підрозділи", None),
                     ("Червоний текст", "що написати в блоці та обсяг у знаках (з пробілами)", None),
                     ("УТОЧНИТИ", "дані, які потрібно отримати від клієнта", None)]:
        S.cell(r, 2, k).font = Font(name=F, size=10, bold=True, color="FFFF0000" if k == "Червоний текст" else None)
        if fl: S.cell(r, 2).fill = fl
        S.cell(r, 3, v).font = f_meta
        r += 1

    A = wb.create_sheet("Аналіз ТОПу", 1)
    A.sheet_view.showGridLines = False
    A["A1"] = f"Аналіз видачі Google Україна ({time.strftime('%d.%m.%Y')})"; A["A1"].font = f_title
    for col, w in zip("ABCDE", [24, 6, 34, 90, 60]):
        A.column_dimensions[col].width = w
    r = 3
    for it in items:
        c = it["tz"]
        A.cell(r, 1, f"{it['id']}. {c.get('name_ua') or c.get('name_ru')}").font = f_h2
        A.cell(r, 1).fill = fill_h2
        A.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        r += 1
        for i, h in enumerate(["Сторінка", "Мова", "Запит", "ТОП-10", "«Люди також питають»"], 1):
            x = A.cell(r, i, h); x.font = f_hdr; x.fill = fill_hdr; x.border = box
        r += 1
        for lang in _langs(it):
            res = it["research"].get(lang, {})
            for s in res.get("serps", []):
                urls = "\n".join(f"{o['position']}. {o['url']}" for o in s["organic"])
                vals = [c.get("short"), lang.upper(), s["query"], urls, " | ".join(s.get("paa", [])) or "—"]
                for i, v in enumerate(vals, 1):
                    x = A.cell(r, i, v); x.font = f_meta; x.alignment = wrap; x.border = box
                A.row_dimensions[r].height = min(409, 13 * max(1, len(s["organic"])) + 4)
                r += 1
            A.cell(r, 3, f"Сторінки конкурентів ({lang.upper()})").font = f_meta_k
            r += 1
            for p in res.get("pages", []):
                A.cell(r, 3, p.get("domain")).font = f_meta
                info = (f"{p['chars']:,} зн. · {p.get('type_guess')} · H2: {sum(1 for h in p.get('headings', []) if h[0] == 'H2')}"
                        .replace(",", " ")) if not p.get("error") else f"не завантажено ({p['error']})"
                A.cell(r, 4, f"{p['url']}  —  {info}").font = f_meta
                r += 1
        A.cell(r, 3, "Висновок").font = f_meta_k
        x = A.cell(r, 4, c.get("top_note", "")); x.font = f_meta; x.alignment = wrap
        A.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        A.row_dimensions[r].height = est_h(c.get("top_note", ""), 140, 13)
        r += 2

    M = wb.create_sheet("Методика", 1)
    M.column_dimensions["A"].width = 120
    M["A1"] = "Як сформовані ТЗ"; M["A1"].font = f_title
    steps = [
        "1. Кожен кластер семантики = одна сторінка. Для мов, де є семантика, — окреме ТЗ (RU / UA).",
        "2. За головними ключами кластера знято видачу Google Україна (UA — hl=uk, RU — hl=ru): ТОП-10, «Люди також питають», пов’язані запити.",
        "3. Сторінки ТОПу (без довідників, бірж і соцмереж) завантажено й розібрано: H1–H3, обсяг основного контенту, FAQ, ціни, входження ключів.",
        "4. Інтент і шаблон визначено за типами сторінок у ТОП: лендинг послуги або гібрид «послуга + експертний контент».",
        "5. Обсяг — з урахуванням медіани ТОПу. Структура — обов’язкові блоки більшості ТОПу + прогалини (позначено «у ТОП цього немає»).",
        "6. LSI-діапазони — за кількістю входжень у ТОПі, приведеною до цільового обсягу. FAQ — з «Люди також питають» і FAQ конкурентів.",
        "7. Title і Description — чернетки, звірити з довжиною сніпета. Дані «УТОЧНИТИ» отримати від клієнта.",
    ]
    for i, s in enumerate(steps, 3):
        x = M.cell(i, 1, s); x.font = f_meta; x.alignment = wrap
    wb.save(path)


# ------------------------------------------------------------------ Base
def build_base(items, path):
    f_head = Font(name=F, size=12, bold=True)
    f_cs = Font(name=F, size=8, color="FFFF0000")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for it in items:
        c = it["tz"]
        for lang in _langs(it):
            L = 0 if lang == "ru" else 1
            ws = wb.create_sheet(safe_title(wb, f"{c['short']} {lang.upper()}"))
            ws.column_dimensions["A"].width = 60
            ws.column_dimensions["B"].width = 5.3
            ws.column_dimensions["C"].width = 38
            ws.column_dimensions["D"].width = 14
            r = 2
            for b in c["blocks"]:
                t = b[0]
                cell = ws.cell(r, 1)
                if t == "h1":
                    cell.value = f"H1 {b[1 + L]}"; cell.font = f_head; cell.fill = fill_h1; cell.alignment = wrap
                    r += 1
                    ws.cell(r, 1, ("объём текста: " if lang == "ru" else "обсяг тексту: ")
                            + str(c.get("volume", "")).replace(" з пробілами", "")).font = f_cs
                elif t == "h2":
                    r += 1
                    cell = ws.cell(r, 1)
                    cell.value = f"H2: {b[1 + L]}"; cell.font = f_head; cell.fill = fill_h2; cell.alignment = wrap
                elif t == "h3":
                    cell.value = f"H3: {b[1 + L]}"; cell.font = f_head; cell.alignment = wrap
                elif t == "c":
                    cell.value = b[1]; cell.font = f_cs; cell.alignment = wrap
                elif t in ("cta", "faq"):
                    r += 1
                    cell = ws.cell(r, 1)
                    cell.value = t.upper(); cell.font = f_head; cell.fill = fill_h2
                elif t in ("q", "t"):
                    cell.value = b[1 + L]; cell.font = Font(name=F, size=10); cell.alignment = wrap
                r += 1
            rr = 2
            for k, f in list(c.get(f"kw_{lang}", [])) + list(c.get(f"lsi_{lang}", [])):
                ws.cell(rr, 3, k).font = Font(name=F, size=10)
                ws.cell(rr, 3).alignment = wrap
                ws.cell(rr, 4, f).font = Font(name=F, size=10)
                rr += 1
    wb.save(path)
