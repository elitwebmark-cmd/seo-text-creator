import importlib
import os
import time
from pathlib import Path

os.environ["MOCK"] = "1"
os.environ["APP_PASSWORD"] = ""
os.environ["DATA_DIR"] = str(Path(__file__).parent / ".data")

from fastapi.testclient import TestClient  # noqa: E402

from app import semantics, analysis  # noqa: E402
from app.main import app  # noqa: E402

client = TestClient(app)


def test_parse_text_and_pair():
    ua = semantics.parse_text("# SERM\nserm послуги; 40\nуправління репутацією в інтернеті; 30")
    ru = semantics.parse_text("# SERM\nserm услуги; 150\nуправление репутацией в интернете; 150")
    pages = semantics.auto_pair(ua, ru)
    assert pages == [{"ua": 0, "ru": 0}]


def test_phrase_count_morphology():
    t = "Послуги маркетолога. Послуг маркетологів немає. послуга маркетолога"
    assert analysis.count_phrase(t, "послуги маркетолога") == 3


def test_template_roundtrip():
    r = client.get("/api/template.xlsx")
    assert r.status_code == 200
    p = semantics.parse_xlsx(r.content)
    assert p["ua"][0]["name"] == "SERM просування" and len(p["ru"][0]["keywords"]) == 2


def test_full_job_mock():
    r = client.post("/api/parse", data={"text_ua": "# SERM\nserm послуги; 40", "text_ru": "# SERM\nserm услуги; 150"})
    assert r.status_code == 200, r.text
    d = r.json()
    body = {"project": {"domain": "elit-web.ua"}, "formats": ["advanced", "base"],
            "pages": [{"name_ua": "SERM", "name_ru": "SERM", "ua": d["ua"][0]["keywords"], "ru": d["ru"][0]["keywords"]}]}
    jid = client.post("/api/jobs", json=body).json()["id"]
    for _ in range(60):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error", "partial"):
            break
        time.sleep(0.5)
    assert j["status"] == "done", j
    assert len(j["files"]) == 2
    for f in j["files"]:
        assert client.get(f"/api/jobs/{jid}/files/{f}").status_code == 200


def test_semantics_job_mock():
    jid = client.post("/api/semantics", json={"entries": ["SERM послуги", "https://example.com/ua/serm"], "langs": ["ua", "ru"]}).json()["id"]
    for _ in range(60):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error", "partial"):
            break
        time.sleep(0.3)
    assert j["status"] == "done", j
    d = client.get(f"/api/jobs/{jid}/semantics").json()
    assert d["ua"] and d["ru"] and d["pages"]
    assert client.get(f"/api/jobs/{jid}/files/Semantics.xlsx").status_code == 200
    kinds = {x["id"]: x["kind"] for x in client.get("/api/jobs").json()}
    assert kinds[jid] == "semantics"


def test_text_without_volume_is_none():
    c = semantics.parse_text("# X\nключ один\nключ два; 30")
    assert dict(c[0]["keywords"]) == {"ключ один": None, "ключ два": 30}


def test_dfs_only_for_missing(monkeypatch):
    import asyncio
    from app import pipeline, volumes, db
    calls = []

    async def fake_sv(kws, lang, log=None):
        calls.append((lang, list(kws)))
        return {k: 70 for k in kws}

    monkeypatch.setattr(volumes, "search_volume", fake_sv)
    monkeypatch.setattr(volumes, "dfs_enabled", lambda: True)
    jid = db.create("t", {})
    full = [{"ua": [("a", 10), ("b", 0)], "ru": [("c", 5)]}]
    asyncio.run(pipeline.fill_missing_volumes(jid, full))
    assert calls == []  # частотність є — DataForSEO не чіпаємо
    part = [{"ua": [("a", 10), ("b", None)], "ru": [("c", 5)]}]
    asyncio.run(pipeline.fill_missing_volumes(jid, part))
    assert calls == [("ua", ["b"])]
    assert dict(part[0]["ua"])["b"] == 70


def test_volume_cache(monkeypatch):
    from app import db
    db.save_volumes({"тест кеш": 40}, "ua")
    assert db.cached_volumes(["тест кеш", "інший"], "ua", 90) == {"тест кеш": 40}


def test_volumes_job_excel_format():
    import openpyxl, io
    body = {"pages": [{"name_ua": "SERM", "name_ru": "SERM", "ua": [["serm послуги", 40]], "ru": [["serm услуги", None]]},
                      {"name_ua": "Локальне SEO", "ua": [["локальне seo", 100]]}]}
    jid = client.post("/api/volumes", json=body).json()["id"]
    for _ in range(40):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error", "partial"):
            break
        time.sleep(0.3)
    assert j["files"] == ["Semantics.xlsx"], j
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/jobs/{jid}/files/Semantics.xlsx").content))
    assert wb.sheetnames == ["ua", "ru"]
    rows = [r for r in wb["ua"].iter_rows(values_only=True)]
    assert rows[1][1] == "SERM" and rows[2][2] == "serm послуги" and rows[3][1] == "Локальне SEO"
    p = semantics.parse_xlsx(client.get(f"/api/jobs/{jid}/files/Semantics.xlsx").content)
    assert [c["name"] for c in p["ua"]] == ["SERM", "Локальне SEO"]


def _wait(jid, n=80):
    for _ in range(n):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error", "partial"):
            return j
        time.sleep(0.3)
    return j


def test_tz_with_texts_and_texts_job():
    import docx, io
    body = {"project": {"domain": "elit-web.ua"}, "formats": ["base", "texts"],
            "pages": [{"name_ua": "SERM", "name_ru": "SERM", "ua": [["serm послуги", 40]], "ru": [["serm услуги", 150]]}]}
    j = _wait(client.post("/api/jobs", json=body).json()["id"])
    assert j["status"] in ("done", "partial"), j
    docs = [f for f in j["files"] if f.endswith(".docx")]
    assert docs, j
    d = docx.Document(io.BytesIO(client.get(f"/api/jobs/{j['id']}/files/{docs[0]}").content))
    assert any(p.style.name.startswith("Heading") for p in d.paragraphs)
    j2 = _wait(client.post(f"/api/jobs/{j['id']}/texts").json()["id"])
    assert any(f.endswith(".docx") for f in j2["files"]), j2


def test_checker_finds_problems():
    from app import writer, prompts, llm
    tz = llm.normalize(prompts.EXAMPLE)
    md = "# Неправильний заголовок\n\nКороткий текст."
    rep = writer.check(md, tz, "ua")
    assert not rep["ok"] and any("обсяг" in x.lower() for x in rep["issues"]) and any("Немає заголовка" in x for x in rep["issues"])


def test_import_existing_tz_and_write_texts():
    import openpyxl, io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SERM RU"
    for i, (a, c, d) in enumerate([("H1 SERM: управление репутацией", "serm услуги", "2-3"),
                                   ("объём текста: 3000–4000 зн", "отзывы", "3-6"),
                                   ("H2: Что такое SERM", None, None), ("400–600 зн. Определение.", None, None),
                                   ("FAQ", None, None), ("Сколько стоит SERM?", None, None)], 1):
        ws.cell(i, 1, a)
        if c:
            ws.cell(i, 3, c); ws.cell(i, 4, d)
    buf = io.BytesIO(); wb.save(buf)
    r = client.post("/api/tz/parse", files={"file": ("tz.xlsx", buf.getvalue())}, data={"text": "H1: Тест\nH2: Розділ один\nH2: Розділ два"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["items"][0]["lang"] == "ru" and d["items"][0]["stats"]["keywords"] == 2
    assert len(d["items"]) == 2  # xlsx-аркуш + вставлений текст
    j = _wait(client.post("/api/tz/texts", json={"token": d["token"], "selected": [0, 1]}).json()["id"])
    assert j["status"] in ("done", "partial") and any(f.endswith(".docx") for f in j["files"]), j


def test_import_example_file_structure():
    from app import tz_import
    from pathlib import Path
    p = Path("/root/.claude/uploads/ab413567-f57b-5cd8-80ea-d81f0a266c86/b7f75787-_______________________________elit-web.ua.xlsx")
    if not p.exists():
        return
    r = tz_import.parse_xlsx(p.read_bytes())
    assert len(r) == 4 and all(x["lang"] == "ru" for x in r)
    assert r[0]["tz"]["blocks"][0][0] == "h1"


def test_styles_catalog_detect_and_brief():
    from app import styles, writer, prompts, llm
    ids = [s["id"] for s in client.get("/api/styles").json()]
    assert ids[:2] == ["default", "auto"] and set("ABVG") <= set(ids)
    assert styles.detect({"name_ua": "SEO-аудит сайту", "kw_ua": [("seo аудит", 50)]}) == "A"
    assert styles.detect({"name_ua": "Налаштування Google Ads", "kw_ua": [("контекстна реклама", 50)]}) == "B"
    assert styles.detect({"name_ua": "SMM просування", "kw_ua": [("просування в instagram", 50)]}) == "V"
    assert styles.detect({"name_ua": "Розробка сайтів", "kw_ua": [("створення сайту під ключ", 50)]}) == "G"
    assert styles.detect({"name_ua": "Щось інше"}) == "default"
    tz = llm.normalize(prompts.EXAMPLE)
    b = writer.tz_brief(tz, "ru", {"facts": "12 років", "media": True}, None, "B")
    assert "ПРАКТИК-РЕЗУЛЬТАТНИК" in b and "РЕДАКЦІЙНІ РАМКИ" in b and "російською" in b and "12 років" not in b
    assert "Стиль і тональність" not in writer.tz_brief(tz, "ua", {}, None, "default")
    md = "# X\n\nУ сучасному світі це не просто сайт, а інструмент. Це не просто текст, а продаж!"
    assert any("AI-штампи" in x for x in writer.check(md, tz, "ua", "A", [], {"text_type": "article"})["issues"])
    assert any("AI-штампи" in x for x in writer.check(md, tz, "ua", "A")["warnings"])  # посадкова — лише рекомендація
    assert any("AI-штампи" in x for x in writer.check(md, tz, "ua")["warnings"])


def test_texts_with_styles_per_page_and_batch():
    import docx, io
    body = {"project": {"domain": "elit-web.ua", "text_style": "auto"}, "formats": ["base", "texts"],
            "pages": [{"name_ua": "SERM", "name_ru": "SERM", "ua": [["serm послуги", 40]], "ru": [["serm услуги", 150]], "style": "G"}]}
    j = _wait(client.post("/api/jobs", json=body).json()["id"])
    doc = [f for f in j["files"] if f.endswith(".docx")][0]
    txt = "\n".join(p.text for p in docx.Document(io.BytesIO(client.get(f"/api/jobs/{j['id']}/files/{doc}").content)).paragraphs)
    assert "Стиль: Г. Радник-стратег" in txt
    pages = client.get(f"/api/jobs/{j['id']}/pages").json()
    assert pages and pages[0]["auto"]
    j2 = _wait(client.post(f"/api/jobs/{j['id']}/texts", json={"style": "A", "styles": {pages[0]["id"]: "B"}, "media": True}).json()["id"])
    import openpyxl
    tech = [f for f in j2["files"] if f.startswith("Tech_")][0]
    ws = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/jobs/{j2['id']}/files/{tech}").content))["Статті"]
    assert "Б. Практик-результатник" in [c.value for c in ws[2]]
    r = client.post("/api/tz/parse", data={"text": "H1: SMM просування\nH2: Розділ один\nH2: Розділ два"}).json()
    j3 = _wait(client.post("/api/tz/texts", json={"token": r["token"], "selected": [0], "style": "auto"}).json()["id"])
    doc = [f for f in j3["files"] if f.endswith(".docx")][0]
    txt = "\n".join(p.text for p in docx.Document(io.BytesIO(client.get(f"/api/jobs/{j3['id']}/files/{doc}").content)).paragraphs)
    assert "В. Голос аудиторії" in txt


def _fact(year=None):
    import datetime as dt
    y = year or dt.date.today().year
    return {"id": "F1", "url": "https://www.thinkwithgoogle.com/s", "org": "Google", "year": y, "geo": "світ", "number": "53%", "fit": 5,
            "quote": f"In {y}, 53% of mobile site visits were abandoned if a page took longer than three seconds to load."}


def test_fact_verification():
    from app import facts
    f = _fact()
    page = {"https://www.thinkwithgoogle.com/s": {"text": "Intro. " + f["quote"].replace("53%", "53 %") + " More text.", "date": ""}}
    assert facts.verify(f, page) == ""
    assert facts.verify({**f, "quote": f["quote"].replace("53%", "63%"), "number": "63%"}, page) != ""
    assert "застарілі" in facts.verify(_fact(2015), {"https://www.thinkwithgoogle.com/s": {"text": _fact(2015)["quote"], "date": ""}})
    assert facts.verify({**f, "url": "https://some-agency.ua/blog"}, {"https://some-agency.ua/blog": page["https://www.thinkwithgoogle.com/s"]}) != ""


def test_checker_numbers_links_placeholders():
    from app import writer, prompts, llm
    tz = llm.normalize(prompts.EXAMPLE)
    f = _fact()
    good = f"# X\n\nЗа даними [Google]({f['url']}), {f['year']}, 53% відвідувань обриваються. Умовно 20% бюджету йде на тести."
    rep = writer.fact_checks(good, tz, "ua", [f], {})
    assert not rep[0] and rep[2] == ["F1"], rep
    bad = ("# X\n\nЗа статистикою, 43% компаній втрачають клієнтів. Ціна: [уточнити: вартість]. "
           "Дивіться [звіт](https://random.blog/x). Google: 53% відвідувань обриваються.")
    issues = writer.fact_checks(bad, tz, "ua", [f], {})[0]
    txt = " | ".join(issues)
    assert "Плейсхолдер" in txt and "не з перевірених" in txt and "без перевіреного джерела" in txt and "без посилання" in txt, txt
    st = writer.structure_checks("# A\n\n## Зовсім новий розділ\n\nтекст", tz, "ua")[0]
    assert any("відсутній у ТЗ" in x for x in st)


def test_editor_regression_keeps_best_version(monkeypatch):
    import asyncio
    from app import writer, prompts, llm, config
    tz = llm.normalize(prompts.EXAMPLE)
    draft = writer._mock_text(tz, "ua")
    first_para = draft.split("\n\n")[1]
    bad_patch = f"<<<<<<< FIND\n{first_para}\n=======\n{first_para}\n\n## Зайвий розділ\n\nЗа статистикою, 43% компаній. [уточнити]\n>>>>>>> END"
    seen = []

    def fake(system, prompt, max_tokens=16000, **kw):
        seen.append(kw.get("step"))
        tail = prompt[-1] if isinstance(prompt, list) else prompt
        if kw.get("step") in ("тексти: автор", "тексти: редактор"):
            assert isinstance(prompt, list) and prompt[0][1] is True  # бриф кешується
        if "шеф-редактор" in tail:
            return bad_patch
        if "Напиши текст" in tail:
            return draft
        return "x" * 300

    async def no_lt(md, lang):
        return []

    monkeypatch.setattr(config, "MOCK", False)
    monkeypatch.setattr(config, "TEXT_AUDIT", False)
    monkeypatch.setattr(llm, "call_text", fake)
    monkeypatch.setattr(writer, "languagetool", no_lt)
    res = asyncio.run(writer.write_text(tz, "ua", {}, None, lambda m: None, facts=[]))
    assert "Зайвий розділ" not in res["markdown"] and res["rounds"] >= 1
    assert "тексти: автор" in seen and "тексти: редактор" in seen


def test_apply_patches():
    from app import writer
    md = "# T\n\nПерше  речення тут.\nДруге речення.\n\n## Розділ\n\nТекст."
    ans = ("<<<<<<< FIND\nПерше речення тут. Друге речення.\n=======\nНове речення.\n>>>>>>> END\n"
           "<<<<<<< FIND\nнемає такого\n=======\nх\n>>>>>>> END")
    new, ok, bad = writer.apply_patches(md, ans)
    assert ok == 1 and bad == 1 and "Нове речення." in new and "Друге" not in new
    assert writer.apply_patches(md, "БЕЗ ЗМІН") == (md, 0, 0)


def test_cost_tracking(monkeypatch):
    from types import SimpleNamespace
    from app import llm, db
    jid = db.create("cost test", {"kind": "tz"})
    tok = llm.JOB.set(jid)
    try:
        llm._track("тексти: автор", "claude-sonnet-4-5", SimpleNamespace(input_tokens=1000, output_tokens=2000,
                   cache_creation_input_tokens=10000, cache_read_input_tokens=0))
        llm._track("тексти: аудит ТЗ", "claude-haiku-4-5", SimpleNamespace(input_tokens=1000, output_tokens=1000,
                   cache_creation_input_tokens=0, cache_read_input_tokens=10000))
    finally:
        llm.JOB.reset(tok)
    c = db.job_costs(jid)
    assert c["calls"] == 2 and abs(c["usd"] - (0.003 + 0.03 + 0.0375 + 0.001 + 0.005 + 0.001)) < 1e-6
    assert any(j["id"] == jid and j["cost_usd"] > 0 for j in db.list_jobs())


def test_docx_has_hyperlink_and_sources():
    import docx, io
    body = {"project": {"domain": "elit-web.ua", "text_type": "article"}, "formats": ["texts"],
            "pages": [{"name_ua": "SERM", "ua": [["serm послуги", 40]]}]}
    j = _wait(client.post("/api/jobs", json=body).json()["id"])
    doc = [f for f in j["files"] if f.endswith(".docx")][0]
    d = docx.Document(io.BytesIO(client.get(f"/api/jobs/{j['id']}/files/{doc}").content))
    xml = d.element.xml
    assert "w:hyperlink" in xml and "thinkwithgoogle" in str([r.target_ref for r in d.part.rels.values()])
    assert any("Факти з джерелами" in p.text for p in d.paragraphs)


def test_text_type_article_format_and_tz_prompt():
    from app import text_types, writer, prompts, llm
    tz = llm.normalize(prompts.EXAMPLE)
    choppy = "# X\n\nКоротко.\n\n## Розділ\n\n- раз\n- два\n- три\n\nЩе.\n\nІ ще.\n"
    rep = writer.check(choppy, tz, "ua", "default", [], {"text_type": "article"})
    assert any("абзаци закороткі" in x for x in rep["issues"]) and any("забагато списків" in x for x in rep["issues"])
    assert not any("Формат статті" in x for x in writer.check(choppy, tz, "ua", "default", [], {})["issues"])
    b = writer.tz_brief(tz, "ua", {"text_type": "article"}, None, "A", [])
    assert "Формат тексту: СТАТТЯ" in b and "Тип тексту: Стаття (блог компанії)" in b
    p = prompts.build_user_prompt({"name_ua": "X", "ua": [("x", 10)]}, {"text_type": "media"}, {})
    assert "ТИП ТЕКСТУ: СТАТТЯ ДЛЯ МЕДІА" in p
    assert "ТИП ТЕКСТУ" not in prompts.build_user_prompt({"name_ua": "X", "ua": [("x", 10)]}, {}, {})
    assert text_types.resolve({"media": True}) == "media" and [t["id"] for t in client.get("/api/text-types").json()] == ["landing", "article", "media"]


def test_texts_job_with_text_type():
    import docx, io
    body = {"project": {"domain": "elit-web.ua", "text_type": "article"}, "formats": ["texts"],
            "pages": [{"name_ua": "SERM", "ua": [["serm послуги", 40]]}]}
    j = _wait(client.post("/api/jobs", json=body).json()["id"])
    doc = [f for f in j["files"] if f.endswith(".docx")][0]
    txt = "\n".join(p.text for p in docx.Document(io.BytesIO(client.get(f"/api/jobs/{j['id']}/files/{doc}").content)).paragraphs)
    assert "Тип: Стаття (блог компанії)" in txt and "Формат:" in txt


def test_media_articles_output_and_tz_adaptation():
    import docx, io, zipfile, openpyxl
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from app import text_types, prompts, llm
    tz = llm.normalize(prompts.EXAMPLE)
    tz["blocks"] = tz["blocks"] + [("h2", "Услуги SERM от агентства: что входит и сколько стоит", "Послуги SERM від агенції: що входить і скільки коштує"),
                                   ("c", "пакети й ціни"), ("h3", "Як працює?", "Як працює?")]
    a = text_types.adapt_tz(tz, "media")
    kinds = [b[0] for b in a["blocks"]]
    assert "cta" not in kinds and kinds[-1] == "q" or kinds.index("faq") > max(i for i, k in enumerate(kinds) if k == "h2")
    assert not any("агенції" in (b[2] if len(b) > 2 else "") for b in a["blocks"])
    assert not any(b[0] in ("h2", "h3") and b[2].endswith("?") for b in a["blocks"])
    assert any("CTA" in d for d in a["_dropped"]) and any("агенції" in d for d in a["_dropped"])
    body = {"project": {"domain": "elit-web.ua", "text_type": "media"}, "formats": ["texts"],
            "pages": [{"name_ua": "SERM", "ua": [["serm послуги", 40]]}]}
    j = _wait(client.post("/api/jobs", json=body).json()["id"])
    z = [f for f in j["files"] if f.endswith(".zip")][0]
    zf = zipfile.ZipFile(io.BytesIO(client.get(f"/api/jobs/{j['id']}/files/{z}").content))
    docs = [n for n in zf.namelist() if n.endswith(".docx")]
    assert docs and any(n.startswith("Tech_") for n in zf.namelist()) and "elit" not in z.lower()
    d = docx.Document(io.BytesIO(zf.read(docs[0])))
    first = next(p for p in d.paragraphs if p.text.strip())
    assert first.style.name == "Heading 1"
    assert all(p.alignment == AL.JUSTIFY for p in d.paragraphs if p.text.strip() and not p.style.name.startswith("Heading"))
    assert not any(p.text.strip().upper() == "CTA" for p in d.paragraphs)
    wb = openpyxl.load_workbook(io.BytesIO(zf.read([n for n in zf.namelist() if n.startswith("Tech_")][0])))
    assert {"Статті", "Ключі", "Джерела"} <= set(wb.sheetnames)


def test_h3_block_as_list_is_accepted_for_articles():
    from app import writer, prompts, llm
    tz = llm.normalize(prompts.EXAMPLE)
    tz["blocks"] = [("h1", "Т", "Т"), ("h2", "Типи", "Типи"), ("h3", "Перший тип?", "Перший тип?"), ("h3", "Другий тип", "Другий тип")]
    md = "# Т\n\n## Типи\n\nВступ.\n\n- **Перший тип.** Пояснення першого.\n- **Другий тип.** Пояснення другого.\n\nПерехід далі."
    rep = writer.check(md, tz, "ua", "default", [], {"text_type": "article"})
    assert not any("Немає заголовка" in x for x in rep["issues"]) and rep["converted"] == ["Типи"], rep["issues"]
    mixed = "# Т\n\n## Типи\n\n### Перший тип\n\nТекст.\n\n- **Другий тип.** Пояснення."
    assert any("змішувати" in x for x in writer.check(mixed, tz, "ua", "default", [], {"text_type": "article"})["issues"])
    assert any("Немає заголовка" in x for x in writer.check(md, tz, "ua", "default", [], {})["issues"])


def test_section_patch_and_truncated_patch():
    from app import writer
    md = "# T\n\nВступ.\n\n## Розділ\n\nСтарий текст.\n\nЩе старий.\n\n## Інший\n\nЛишається."
    ans = ("<<<<<<< SECTION\n## Розділ\nНовий розгорнутий текст розділу.\n>>>>>>> END\n"
           "<<<<<<< FIND\nЛишається.\n=======\nОбріз")  # остання правка обрізана — ігнорується
    new, ok, bad = writer.apply_patches(md, ans)
    assert ok == 1 and "Новий розгорнутий текст розділу." in new and "Старий" not in new and "Лишається." in new and "## Інший" in new


def test_call_text_continues_after_max_tokens(monkeypatch):
    from types import SimpleNamespace
    from app import llm, config
    calls = []

    class Stream:
        def __init__(self, **kw):
            calls.append(kw["messages"])
        def __enter__(self):
            return self
        def __exit__(self, *a):
            return False
        def get_final_message(self):
            n = len(calls)
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="Частина один. " * 20 if n == 1 else "Частина два.")],
                                   stop_reason="max_tokens" if n == 1 else "end_turn",
                                   usage=SimpleNamespace(input_tokens=1, output_tokens=1, cache_creation_input_tokens=0, cache_read_input_tokens=0))

    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "x")
    monkeypatch.setattr(llm, "_client", lambda: SimpleNamespace(messages=SimpleNamespace(stream=lambda **kw: Stream(**kw))))
    out = llm.call_text("sys", "prompt", 100, step="тест")
    assert out.endswith("Частина два.") and len(calls) == 2 and calls[1][-1]["role"] == "assistant"


def test_landing_is_simple_and_editor_gated(monkeypatch):
    import asyncio
    from app import writer, prompts, llm, config, text_types
    tz = llm.normalize(prompts.EXAMPLE)
    # дрібні проблеми (LSI, FAQ, перевищення ключа на 1) не кличуть редактора
    rep = {"issues": ["Відповідь FAQ №1: 230 зн., треба 250–400", "Розділ «X»: 900 зн., ціль 1000–1200"], "chars": 9000}
    assert writer.blocking(rep, tz) == []
    assert writer.blocking({"issues": ["Немає заголовка (або порушено порядок): «X»"], "chars": 9000}, tz)
    # посадкова: цифра без джерела — лише попередження
    r = writer.check("# A\n\nЗа даними досліджень, 43% компаній так роблять.", tz, "ua", "default", None, {})
    assert not any("Цифра без перевіреного" in x for x in r["issues"]) and any("Цифра без перевіреного" in x for x in r["warnings"])
    # автофікс для статей
    fixed = text_types and writer.autofix("# T\n\n## Що таке X?\n\nТак, **важливо** це.\n\n## FAQ\n\n### Скільки коштує\n\nВідповідь.", "media")
    assert "## Що таке X\n" in fixed and "**" not in fixed and "Важливо це." in fixed and "### Скільки коштує?" in fixed
    # посадкова з чистою чернеткою — без редактора, без коректури, без аудиту
    good = writer._mock_text(tz, "ua")
    steps = []

    def fake(system, prompt, max_tokens=16000, **kw):
        steps.append(kw.get("step"))
        return good

    async def no_lt(md, lang):
        return []

    monkeypatch.setattr(config, "MOCK", False)
    monkeypatch.setattr(llm, "call_text", fake)
    monkeypatch.setattr(writer, "languagetool", no_lt)
    monkeypatch.setattr(writer, "blocking", lambda rep, tz: [])
    asyncio.run(writer.write_text(tz, "ua", {}, None, lambda m: None))
    assert steps == ["тексти: автор"]


def test_volumes_ru_fallback_and_bad_keyword(monkeypatch):
    import asyncio
    from app import volumes, config
    sent = []

    async def fake_dfs(endpoint, payload):
        sent.append(payload)
        if "language_code" in payload:  # мову не прийнято для цієї локації
            raise RuntimeError("40501 Invalid Field: 'language_code'")
        if any("погано" in k for k in payload["keywords"]):
            raise RuntimeError("40501 Invalid keyword")
        return [{"keyword": k.replace("е", "ё", 1) if k.startswith("реклама") else k, "search_volume": 50} for k in payload["keywords"]]

    monkeypatch.setattr(volumes, "_dfs", fake_dfs)
    monkeypatch.setattr(volumes, "dfs_enabled", lambda: True)
    monkeypatch.setattr(config, "MOCK", False)
    logs = []
    out = asyncio.run(volumes.search_volume(["Реклама в Telegram Ads?", "telegram ads цена", "погано"], "ru", logs.append))
    assert out["Реклама в Telegram Ads?"] == 50 and out["telegram ads цена"] == 50 and out.get("погано", 0) == 0
    assert any("без прив'язки до мови" in l for l in logs)
